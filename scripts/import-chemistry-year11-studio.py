from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E,html
from docx import Document
import re,json,copy,subprocess,hashlib,importlib.util,html as escaping
import fitz
import io,base64
from PIL import Image
ROOT=Path(__file__).resolve().parents[1];WORK=ROOT/'chemistry-year11-work';OUT=WORK/'output';OUT.mkdir(exist_ok=True,parents=True)
helper_path=ROOT/'physics-year11-work/import_year11.py'
if not helper_path.exists():helper_path=ROOT/'scripts/import-year11-physics-studio.py'
(ROOT/'physics-year11-work').mkdir(exist_ok=True)
spec=importlib.util.spec_from_file_location('physics_tools',helper_path);helpers=importlib.util.module_from_spec(spec);spec.loader.exec_module(helpers);helpers.WORK=WORK
original_content=helpers.question_html

def native_image(src):
 if not src.startswith('data:image/'):return src
 payload=base64.b64decode(src.split(',')[1])
 if re.match(r'data:image/(?:x-)?(?:emf|wmf)',src):
  start=payload.find(b'%PDF-');end=payload.rfind(b'%%EOF')
  if start>=0 and end>start:
   with fitz.open(stream=payload[start:end+5],filetype='pdf') as doc:
    pix=doc[0].get_pixmap(matrix=fitz.Matrix(2.5,2.5),alpha=False)
    image=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
  else:
   # Preserve legacy WMF figures using the existing standalone metafile renderer.
   node=html.Element('p');E.SubElement(node,'img',src=src)
   rendered,_=original_content({'nodes':[node]},None,[])
   return html.fragment_fromstring(rendered,create_parent='div').xpath('.//img')[0].get('src')
 else:image=Image.open(io.BytesIO(payload)).convert('RGB')
 image.thumbnail((1800,1800));out=io.BytesIO();image.save(out,format='WEBP',quality=94,method=6)
 return 'data:image/webp;base64,'+base64.b64encode(out.getvalue()).decode()

def chemistry_content(q,doc,pieces,is_key=False):
 if doc is None:return original_content(q,None,[],is_key)
 nodes=copy.deepcopy(q['nodes']);figures=[]
 if q.get('nativeOnly'):pieces=[]
 for node in nodes:
  for image in node.xpath('.//img'):
   if re.match(r'data:image/(?:x-)?(?:emf|wmf)',image.get('src','')):
    image.set('src',native_image(image.get('src')));image.set('style','max-width:100%;height:auto');image.set('alt','Original chemical structure or diagram')
   else:image.getparent().remove(image)
 for pi,clip in pieces:
  page=doc[pi]
  for rect in helpers.figure_regions(page,clip):
   # Keep the diagram crop clear of prose immediately following the illustration.
   graphic_bottom=max([d['rect'].y1 for d in page.get_drawings() if rect.contains(d['rect'])]+[fitz.Rect(im['bbox']).y1 for im in page.get_image_info() if rect.contains(fitz.Rect(im['bbox']))]+[rect.y0])
   for wording,box in helpers.lines(page):
    if box.y0>=graphic_bottom and re.match(r'^(Which|What|How|State|Identify|The absorbance|A student|A water)',wording):
     rect.y1=min(rect.y1,box.y0-2)

   graphics=[d for d in page.get_drawings() if rect.contains(d['rect']) and not(d.get('color') is None and (d.get('fill') is None or min(d['fill'])>.97))]
   if graphics:
    bounds=fitz.Rect(graphics[0]['rect'])
    for g in graphics[1:]:bounds|=g['rect']
    if bounds.width<30 and bounds.height<30:continue
   coloured=any(d.get('fill') and min(d['fill'])<.9 for d in graphics) or any(any(item[0]=='c' for item in d['items']) for d in graphics)
   words=set(re.findall(r'\w+',clean(page.get_text(clip=rect)).lower()))
   if not coloured and len(words)>=4 and any(len(words&set(re.findall(r'\w+',text(n).lower())))/len(words)>.8 for n in nodes if n.tag=='table'):continue
   fig=html.Element('figure');E.SubElement(fig,'img',src=helpers.image_data(page,rect),alt=('Marking guide' if is_key else 'Exam')+' diagram for question '+str(q['number']),loading='lazy',style='max-width:100%;height:auto');figures.append((pi,rect,fig))
 for pi,rect,fig in reversed(figures):
  previous=[(t,b) for t,b in helpers.lines(doc[pi]) if b.y1<rect.y0+4 and len(t)>20 and not re.match(r'Question|See [Nn]ext|Chemistry',t,re.I)];anchor=previous[-1][0].lower() if previous else '';position=0
  for i,node in enumerate(nodes):
   if anchor and(anchor in text(node).lower() or anchor[:45] in text(node).lower()):position=i+1
  nodes.insert(position,fig)
 div=html.Element('div',{'class':'source-text-question'})
 for node in nodes:div.append(node)
 return html.tostring(div,encoding='unicode'),len(figures)
helpers.question_html=chemistry_content
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';NS={'w':W}
PAIRS=[('2020', '1', '2020_chemistry_unit_1_exam.docx', '2020_chemistry_unit_1_marking_guide_final (2).docx'), ('2020', '12', '2020_chemistry_units_1_2.docx', '2020_chemistry_units_1_2_marking_guide.docx'), ('2021', '1', '2021_chemistry_unit_1_exam(1).docx', '2021_chemistry_unit_1_marking_guide(1).docx'), ('2021', '12', '2021_chemistry_units_12_exam(1).docx', '2021_chemistry_units_12_marking_guide(1).docx'), ('2022', '1', '2022_chemistry_unit_1_exam_.docx', '2022_chemistry_unit_1_marking_guide.docx'), ('2022', '12', '2022_chemistry_units_12_exam.docx', '2022_chemistry_units_12_marking_guide.docx'), ('2023', '1', '2023_chemistry_unit_1_exam(1).docx', '2023_chemistry_unit_1_marking_guide(1).docx'), ('2023', '12', '2023_chemistry_units_12_exam(1).docx', '2023_chemistry_units_12_marking_guide(1).docx'), ('2024', '1', '2024_chemistry_unit_1_exam(1).docx', '2024_chemistry_unit_1_marking_guide(1).docx')]

def clean(s):return ' '.join(s.split())
def text(n):
 clone=copy.deepcopy(n)
 for a in clone.xpath('.//annotation'):a.getparent().remove(a)
 return clean(' '.join(clone.itertext()))
def marker(s):
 p=E.Element('{'+W+'}p');r=E.SubElement(p,'{'+W+'}r');E.SubElement(r,'{'+W+'}t').text=s;return p
def section_of(t):
 if re.match(r'This section has 25 questions',t,re.I):return 'Multiple choice'
 if len(t)>160:return None
 if re.match(r'Section\s+(One|1)\s*:\s*Multiple',t,re.I):return 'Multiple choice'
 if re.match(r'Section\s+(Two|2)\s*:\s*Short',t,re.I):return 'Short answer'
 if re.match(r'Section\s+(Three|3)\s*:\s*Extended',t,re.I):return 'Extended answer'
def context_numbers(t):
 m=re.search(r'\bQuestions?\s+([\d\s,–—\-&andto]+)',t,re.I)
 if not m or not re.search(r'refer|answer|use|below|above|following',t,re.I):return []
 part=m[1];nums=[int(n) for n in re.findall(r'\d+',part)];
 if len(nums)==2 and re.search(r'–|—|-|\bto\b',part):nums=list(range(nums[0],nums[1]+1))
 return [n for n in nums if 1<=n<=25]
def preserve_symbols(root):
 from reportlab.pdfbase import _fontdata
 from fontTools.agl import toUnicode
 encoding=_fontdata.encodings['SymbolEncoding'];ns='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
 # Convert Symbol-font text before materialising standalone w:sym elements.
 for run in root.findall('.//{'+ns+'}r'):
  fonts=run.find('{'+ns+'}rPr/{'+ns+'}rFonts')
  if fonts is not None and any(fonts.get('{'+ns+'}'+name)=='Symbol' for name in ['ascii','hAnsi']):
   for t in run.findall('{'+ns+'}t'):
    t.text=''.join((toUnicode(encoding[ord(c)]) if ord(c)<256 and encoding[ord(c)] else c) for c in (t.text or ''))
   for k in list(fonts.attrib):fonts.set(k,'Arial')
 for sym in root.findall('.//{'+ns+'}sym'):
  if sym.get('{'+ns+'}font')=='Symbol':
   code=int(sym.get('{'+ns+'}char'),16)&255;name=encoding[code]
   sym.tag='{'+ns+'}t';sym.attrib.clear();sym.text=toUnicode(name)

def normalise(f):
 folder=WORK/'normalized';folder.mkdir(exist_ok=True);target=folder/f.name;headings=[];highlight={};labels={}
 with ZipFile(f) as z:
  root=E.fromstring(z.read('word/document.xml'));preserve_symbols(root);body=root.find('w:body',NS);kind=None;mcq=0;current=None;option_count=0;formats={}
  if 'word/numbering.xml' in z.namelist():
   numbering=E.fromstring(z.read('word/numbering.xml'))
   for num in numbering.findall('w:num',NS):
    ident=num.get('{'+W+'}numId');ref=num.find('w:abstractNumId',NS).get('{'+W+'}val');abstract=numbering.find('w:abstractNum[@w:abstractNumId="'+ref+'"]',NS)
    if abstract is not None:
     for lvl in abstract.findall('w:lvl',NS):
      fmt=lvl.find('w:numFmt',NS)
      if fmt is not None:formats[(ident,lvl.get('{'+W+'}ilvl'))]=fmt.get('{'+W+'}val')
  for node in list(body):
   if node.tag!='{'+W+'}p':continue
   raw=clean(''.join((c.text or '') if c.tag=='{'+W+'}t' else ' ' if c.tag in ['{'+W+'}tab','{'+W+'}br'] else '' for c in node.xpath('./w:r/*',namespaces=NS)));sec=section_of(raw)
   if sec:kind=sec;body.insert(body.index(node),marker('LVSECTION '+sec));continue
   if kind is None:continue
   if re.match(r'End of (Questions|Examination)|Acknowledg',raw,re.I):body.insert(body.index(node),marker('LVSTOP'));kind=None;continue
   if re.match(r'End of Section|Question\s+\d+\s*\(?continued',raw,re.I):continue
   if kind=='Multiple choice':
    context=context_numbers(raw)
    if context:body.insert(body.index(node),marker('LVCONTEXT '+','.join(map(str,context))))
    num=node.find('w:pPr/w:numPr',NS);level=num.find('w:ilvl',NS).get('{'+W+'}val') if num is not None and num.find('w:ilvl',NS) is not None else None
    number_id=num.find('w:numId',NS).get('{'+W+'}val') if num is not None and num.find('w:numId',NS) is not None else None;fmt=formats.get((number_id,level))
    explicit=re.match(r'^(\d+)(?:\.\s*|\s+)(.*)',raw)
    is_question=bool(explicit and int(explicit[1])==mcq+1)
    if is_question:
     mcq+=1;current=mcq;option_count=0;body.insert(body.index(node),marker('LVQUESTION '+str(mcq)+'|Multiple choice|1'));headings.append(mcq)
     if num is not None:num.getparent().remove(num)
     if explicit:
      # Remove the printed number from its first text run(s), leaving the stem.
      rem=re.match(r'^\s*\d+(?:\.\s*|\s+)',raw)[0];remaining=len(rem)
      ts=node.xpath('./w:r/w:t',namespaces=NS)
      if ts:
       whole=''.join(t.text or '' for t in ts);whole=re.sub(r'^\s*\d+\.\s*','',whole)
       # Tabs are independent OOXML nodes; stripping only the number avoids text loss.
       ts[0].text=re.sub(r'^\s*\d+\.\s*','',ts[0].text or '')
     continue
    choice=re.match(r'^\(?([a-d])\)\s*',raw,re.I)
    if num is not None and fmt in ['lowerLetter','upperLetter']:
     letter=chr(97+option_count);choice_letter=letter;option_count+=1;num.getparent().remove(num)
     r=E.Element('{'+W+'}r');E.SubElement(r,'{'+W+'}t').text='('+letter+') ';node.insert(1,r)
    else:choice_letter=choice[1].lower() if choice else None
    if choice_letter and current:
     marked=node.xpath('.//w:r[w:rPr/w:highlight[@w:val="yellow"] or w:rPr/w:color[@w:val="FF0000"]]/w:t/text()',namespaces=NS)
     if any(s.strip() for s in marked):highlight.setdefault(current,[]).append(choice_letter.upper())
   else:
    match=re.fullmatch(r'Question\s+(\d+)\s*\((\d+)\s*marks?\)',raw,re.I)
    if match:
     current=int(match[1]);body.insert(body.index(node),marker('LVQUESTION '+match[1]+'|'+kind+'|'+match[2]));headings.append(current)
     # Heading is replaced by the metadata marker.
     body.remove(node);continue
   if current:
    for box in node.xpath('.//w:txbxContent/w:p',namespaces=NS):
     t=clean(' '.join(box.xpath('.//w:t/text()',namespaces=NS)))
     if t and t not in labels.setdefault(current,[]):labels[current].append(t)
  with ZipFile(target,'w',ZIP_DEFLATED) as dest:
   for n in z.namelist():dest.writestr(n,E.tostring(root) if n=='word/document.xml' else z.read(n))
 dst=target.with_suffix('.html');subprocess.run(['pandoc',str(target),'-t','html5','--mathml','--embed-resources','-s','--metadata','title=Chemistry','-o',str(dst)],check=True,capture_output=True)
 return dst,highlight,labels,headings
def split_source(f):
 converted,highlight,labels,headings=normalise(f);body=html.parse(str(converted),parser=html.HTMLParser(huge_tree=True)).find('body');qs={};current=None;context=None;contexts=[];kind=None
 for node in body:
  t=text(node)
  if t.startswith('LVSECTION '):kind=t[10:];current=None;context=None;continue
  if t=='LVSTOP':current=None;context=None;kind=None;continue
  if kind is None:continue
  if t.startswith('LVCONTEXT '):context={'numbers':[int(x) for x in t[10:].split(',')],'nodes':[]};contexts.append(context);current=None;continue
  if t.startswith('LVQUESTION '):
   num,category,marks=t[11:].split('|');num=int(num);current={'number':num,'type':category,'marks':int(marks),'nodes':[],'choices':[],'prefix':[]};qs[num]=current;context=None;continue
  if not t and not node.xpath('.//img|.//math'):continue
  if re.match(r'End of Section|Section\s|This section|Suggested working|Question\s+\d+\s*\(?continued|This page.*blank',t,re.I) or re.fullmatch(r'[_\s.]+',t):continue
  if context is not None:context['nodes'].append(copy.deepcopy(node));continue
  if current is None:continue
  if current['type']=='Multiple choice':
   m=re.match(r'^\(([a-d])\)\s*(.*)',t,re.I) if node.tag!='table' else None
   if m and len(re.findall(r'\([a-d]\)',t,re.I))>1:
    chunks=re.split(r'\(([a-d])\)\s*',t,flags=re.I)
    for j in range(1,len(chunks),2):
     piece=html.Element('p');piece.text='('+chunks[j]+') '+chunks[j+1]
     current['choices'].append({'letter':chunks[j].upper(),'text':chunks[j+1].strip(),'node':piece})
    continue
   if m:
    current['choices'].append({'letter':m[1].upper(),'text':m[2],'node':copy.deepcopy(node)});continue
   if node.tag=='table':
    option_cells=[cell for cell in node.xpath('.//td|.//th') if re.match(r'^\([a-d]\)',text(cell),re.I)]
    if len(option_cells)==4:
     for cell in option_cells:
      cm=re.match(r'^\(([a-d])\)\s*(.*)',text(cell),re.I)
      current['choices'].append({'letter':cm[1].upper(),'text':cm[2],'node':copy.deepcopy(cell)})
     current['nodes'].append(copy.deepcopy(node));continue
    option_rows=node.xpath('.//tr[td//ol[@type="a"]]')
    if len(option_rows)==4:
     headings=[text(cell) for cell in node.xpath('.//thead/tr/*') if text(cell)]
     for i,row in enumerate(option_rows):
      for ol in row.xpath('.//ol'):ol.getparent().remove(ol)
      row[0].text='('+chr(97+i)+')'
      cells=[text(cell) for cell in list(row)[2:]]
      wording='; '.join((headings[j]+': ' if j<len(headings) else '')+v for j,v in enumerate(cells))
      current['choices'].append({'letter':chr(65+i),'text':wording,'node':copy.deepcopy(row)})
     current['nodes'].append(copy.deepcopy(node));continue
    rows=node.xpath('.//tr');options=[]
    for row in rows:
     m=re.match(r'^\(([a-d])\)\s*(.*)',text(row),re.I)
     if m:options.append({'letter':m[1].upper(),'text':m[2],'node':copy.deepcopy(row)})
    if len(options)==4:current['choices']+=options;current['nodes'].append(copy.deepcopy(node));continue
   if len(current['choices'])==4:
    # Material after option (d) introduces the following question or question group.
    context={'numbers':[],'after':current['number'],'nodes':[copy.deepcopy(node)]};contexts.append(context);current=None;continue
  current['nodes'].append(copy.deepcopy(node))
 for c in contexts:
  nums=c['numbers'] or [c['after']+1]
  if not c['numbers']:
   words=lambda s:{w.rstrip('s') for w in re.findall(r'[a-z]{5,}',s.lower()) if w not in {'question','following','below','above','would','which','these','their','shown','information'}}
   source_words=words(' '.join(text(n) for n in c['nodes']))
   next_explicit=min([min(other['numbers']) for other in contexts if other['numbers'] and min(other['numbers'])>nums[0]] or [26])
   for following in range(nums[0]+1,min(nums[0]+4,next_explicit)):
    if following not in qs:continue
    stem=' '.join(text(n) for n in qs[following]['nodes']);overlap=source_words&words(stem)
    if len(overlap)>=2 or (overlap and re.search('graph|diagram|image|experiment|participant|structure|shown',stem,re.I)):nums.append(following)
  for n in nums:
   if n in qs:qs[n]['prefix']+=copy.deepcopy(c['nodes'])
 return qs,highlight,labels,contexts
def key_answers(f,highlights):
 answers={}
 for table in Document(f).tables:
  for row in table.rows:
   cells=[c.text.strip() for c in row.cells]
   for i in range(len(cells)-1):
    square=re.findall(r'([a-d])\s*■',cells[i+1],re.I)
    if re.fullmatch(r'\d+',cells[i]) and len(square)==1 and 1<=int(cells[i])<=25:answers[int(cells[i])]=square[0].upper()
    if re.fullmatch(r'\d+',cells[i]) and re.fullmatch('[A-Da-d]',cells[i+1]) and 1<=int(cells[i])<=25:answers[int(cells[i])]=cells[i+1].upper()
 # Answers indicated by colour in table-based alternatives.
 current=None
 normalized=Document(WORK/'normalized'/f.name)
 for node in normalized.element.body:
  if E.QName(node).localname not in ['p','tbl']:continue
  wording=''.join(node.xpath('.//w:t/text()'))
  marker_match=re.match(r'LVQUESTION (\d+)\|',wording)
  if marker_match:current=int(marker_match[1])
  if E.QName(node).localname=='tbl' and current and current<=25:
   for row in node.xpath('./w:tr'):
    wording=' '.join(row.xpath('.//w:t/text()'));choice=re.match(r'^\(([a-d])\)',wording,re.I)
    marked=row.xpath('.//w:r[w:rPr/w:color[@w:val="FF0000"] or w:rPr/w:highlight[@w:val="yellow"]]')
    if choice and marked:highlights.setdefault(current,[]).append(choice[1].upper())
 for n,letters in highlights.items():
  letters=set(letters)
  if len(letters)==1:
   value=next(iter(letters))
   if n in answers and answers[n]!=value:raise ValueError((f.name,n,'conflicting answer key'))
   answers[n]=value
 return answers
def pdf_ranges(f,qs):
 doc=fitz.open(WORK/'repaired-pdf'/(f.stem+'.pdf'));starts={};kind=None;last=0;stop=None;contexts=[];section_start=None
 prefix_groups={}
 for n,q in qs.items():
  first=next((text(node) for node in q['prefix'] if len(text(node))>25),'')
  if first and not context_numbers(first):prefix_groups.setdefault(first.lower()[:35],[]).append(n)
 for pi,p in enumerate(doc):
  ls=helpers.lines(p)
  for t,b in ls:
   sec=section_of(t)
   if sec and (re.match(r'Section\s+(?:One|Two|Three|1|2|3)\s*:',t,re.I) or t.startswith('This section has 25 questions')):
    kind=sec
    if sec=='Multiple choice':section_start=(pi,b.y1+3)
   nums=context_numbers(t)
   if nums and kind=='Multiple choice':contexts.append({'numbers':nums,'start':(pi,b.y0)})
   if kind=='Multiple choice':
    for prefix,numbers in prefix_groups.items():
     if t.lower().startswith(prefix):contexts.append({'numbers':numbers,'start':(pi,b.y0)})
   m=re.match(r'^(\d+)(?:\.\s*|\s+|$)(.*)',t) if kind=='Multiple choice' else re.match(r'Question\s+(\d+)\b(?!\s*\(?continued)',t,re.I)
   if kind=='Multiple choice' and t.isdigit() and (b.y0<65 or b.x0>150):m=None
   if kind=='Multiple choice' and m and int(m[1])>25:m=None
   if kind=='Multiple choice' and m and int(m[1])!=last+1:m=None
   if m and int(m[1]) in qs and int(m[1]) not in starts and int(m[1])>last:
    n=int(m[1]);starts[n]=(pi,b.y0,b.y1+3);last=n
   if last>=max(qs) and re.match(r'^(End of Questions|End of Examination|Acknowledg|Supplementary page$)',t,re.I):stop=(pi,b.y0);break
  if stop:break
 if not stop:stop=(len(doc)-1,doc[-1].rect.height-52)
 result={};numbers=sorted(starts)
 for ix,n in enumerate(numbers):
  pi,top,bodytop=starts[n];ep,ey=starts[numbers[ix+1]][:2] if ix+1<len(numbers) else stop
  if qs[n]['type']!='Multiple choice':top=bodytop
  chunks=[]
  for p in range(pi,ep+1):
   y=top if p==pi else 62;bottom=ey-4 if p==ep else doc[p].rect.height-50
   if bottom>y:chunks.append((p,fitz.Rect(30,y,doc[p].rect.width-30,bottom)))
  result[n]=chunks
 for n,chunks in result.items():
  if qs[n]['type']!='Multiple choice':continue
  cropped=[]
  for pi,rect in chunks:
   boundaries=[c['start'][1]-4 for c in contexts if c['start'][0]==pi and n not in c['numbers'] and min(c['numbers'])>n]
   if boundaries:rect.y1=min(rect.y1,min(boundaries))
   if rect.height>0:cropped.append((pi,rect))
  result[n]=cropped
 # Add shared diagrams placed before the first question in their group.
 for c in contexts:
  first=min(c['numbers']);target=next((starts[n] for n in sorted(starts) if n>=first),None)
  if target is None:continue
  sp,sy=c['start'];ep,ey=target[:2]
  chunks=[]
  for p in range(sp,ep+1):
   top=sy if p==sp else 62;bottom=ey-4 if p==ep else doc[p].rect.height-50
   if bottom>top:chunks.append((p,fitz.Rect(30,top,doc[p].rect.width-30,bottom)))
  for n in c['numbers']:
   if n in result:result[n]=chunks+result[n]
 return doc,result,starts
PATTERNS=[
('Atomic structure and periodicity','Subatomic particles isotopes and ions',r'proton|neutron|isotop|mass number|atomic number|mass spectr|subatomic'),
('Atomic structure and periodicity','Electron configuration and periodic trends',r'electron configuration|electron shell|valence|periodic table|ionisation|ionization|electronegat|atomic radi|group \d|period \d|orbit'),
('Chemical bonding and structure','Ionic metallic and covalent bonding',r'ionic|covalent|metallic|delocalis|delocaliz|bond|lattice'),
('Chemical bonding and structure','Molecular shape polarity and intermolecular forces',r'lewis|electron pair|molecular shape|polarity|polar |dipole|hydrogen bond|dispersion|intermolecular|vapour pressure'),
('Chemical bonding and structure','Allotropes nanomaterials and material properties',r'allotrop|nano|diamond|graphite|graphene|fullerene|polymer|material|conduct|brittle|hardness'),
('Chemical reactions','Formulae nomenclature and equations',r'formula|formulae|nomenclature|name the|iupac|equation|balance|ionic compound'),
('Chemical reactions','Reaction types and observations',r'precipitat|reactivity|reactive metal|observation|redox|oxidation|reduction|displacement|hydrolysis'),
('Chemical calculations','Moles mass and composition',r'mole|mol\b|molar|molarity|avogadro|percent.*composition|percentage.*composition|empirical|molecular formula|relative atomic mass|relative molecular mass|mass of|mass in grams'),
('Chemical calculations','Stoichiometry limiting reagents and yield',r'stoichiometr|limiting|excess|yield|react.*completely|complete reaction|amount.*react|theoretical'),
('Organic chemistry','Hydrocarbons isomers and nomenclature',r'organic|hydrocarbon|alkane|alkene|alkyne|benzene|isomer|structural formula|iupac|methyl|ethyl|propyl|hexane|heptane|octane|pentane|ethene'),
('Organic chemistry','Organic reactions and fuels',r'addition reaction|substitution|polymeris|polymeriz|bromine|combustion|fuel|biofuel|bioethanol|crude oil|petroleum|biogas'),
('Energy and rates','Enthalpy thermochemistry and energy profiles',r'enthalpy|endotherm|exotherm|energy profile|activation energy|heat of|calorimet|specific heat|kilojoule|heat.*released'),
('Energy and rates','Reaction rates and collision theory',r'reaction rate|rate of reaction|collision|catalyst|catalys|surface area|reaction speed'),
('States of matter','Kinetic theory and gas laws',r'gas law|pressure|volume of.*gas|gas.*volume|boyle|charles|ideal gas|kinetic|particle.*motion|temperature.*gas'),
('States of matter','Phase changes and physical properties',r'melt|boil|freez|solidif|vaporis|vaporiz|condens|phase change|phase diagram|heating curve|cooling curve|sublim'),
('Solutions and acids','Solutions solubility and concentration',r'solubil|soluble|solution|concentration|dilut|saturat|dissolv|solvent|solut'),
('Solutions and acids','Acids bases and indicators',r'\bacid|\bbase|\bbasic|\balkali|\bph\b|neutralis|neutraliz|indicator|titration|hydronium|hydroxide'),
('Science inquiry','Mixtures separation and analysis',r'mixture|pure substance|separat|filtrat|distillat|chromatograph|decant|evaporation|spectroscop|spectrophotomet|flame test|absorbance'),
('Science inquiry','Experimental design data and measurement',r'accuracy|accurate|precision|precise|reliab|valid|systematic|random error|uncertainty|experiment|investigat|independent variable|dependent variable|graph|hypothesis')]
def classify(q):
 t=' '.join(text(n) for n in q['prefix']+q['nodes']).lower();matches=[(topic,sub) for topic,sub,pattern in PATTERNS if re.search(pattern,t)]
 if not matches:matches=[('Science inquiry','Experimental design data and measurement')]
 return list(dict.fromkeys(x[0] for x in matches)),list(dict.fromkeys(x[1] for x in matches))
def choice_html(choice):
 clone=copy.deepcopy(choice['node']);clone.tag='div'
 for im in clone.xpath('.//img'):
  im.set('src',native_image(im.get('src','')));im.set('style','max-width:100%;height:auto');im.set('alt','Structure for option '+choice['letter'])
 for node in clone.iter():
  if node.text and node.text.strip():
   node.text=re.sub(r'^\s*\([a-d]\)\s*','',node.text,flags=re.I);break
 for node in list(clone.xpath('.//table|.//tr|.//td|.//th|.//thead|.//tbody')):node.tag='span'
 return html.tostring(clone,encoding='unicode') if text(clone) or clone.xpath('.//img') else '<span>Option '+choice['letter']+' (see source diagram)</span>'
def signature(q):
 chunks=[]
 for root in copy.deepcopy(q['prefix']+q['nodes']+[c['node'] for c in q['choices']]):
  for node in root.xpath('.//annotation'):node.getparent().remove(node)
  def walk(node):
   part=node.text or ''
   for child in node:
    value=walk(child)
    if child.tag=='sub':value='_('+value+')'
    if child.tag=='sup':value='^('+value+')'
    if child.tag=='img':value=' IMAGE '+hashlib.sha256(child.get('src','').encode()).hexdigest()
    part+=value+(child.tail or '')
   return part
  chunks.append(walk(root))
 return re.sub(r'[_\s]+',' ',' '.join(chunks)).strip().casefold()
def main():
 report={'exams':[],'held':[],'duplicates':[],'corrections':[]};records=[];seen={};audits=[]
 for year,units,paper,key in PAIRS:
  f=ROOT/'upload'/paper;kf=ROOT/'upload'/key
  qs,_,labels,contexts=split_source(f);ks,highlight,_,_=split_source(kf);answers=key_answers(kf,highlight)
  assert set(answers)==set(range(1,26)),(key,answers)
  assert sorted(qs)==list(range(1,max(qs)+1)),(paper,sorted(qs))
  pdf,ranges,starts=pdf_ranges(f,qs);kpdf,kranges,_=pdf_ranges(kf,ks)
  assert not(set(qs)-set(ranges)),(paper,'PDF boundaries missing',set(qs)-set(ranges))
  batch=[];filename='practice-studio-chemistry-'+year+'-u'+units+'.json'
  for n,q in sorted(qs.items()):
   key_number=n;options=[];options_html=[];correct=None
   if q['type']=='Multiple choice':
    if sorted(c['letter'] for c in q['choices'])!=['A','B','C','D']:
     report['held'].append({'source':paper,'number':n,'reason':'Four answer choices could not be extracted'});continue
    choices=sorted(q['choices'],key=lambda c:c['letter']);options=[c['text'] or 'Option '+c['letter']+' (see source diagram)' for c in choices];options_html=[choice_html(c) for c in choices]
    correct=ord(answers[n])-65;answer='<p>Correct answer: <strong>'+answers[n]+'</strong></p>'+options_html[correct]+'<p>Answer taken from the supplied exam marking guide.</p>'
   else:
    if year=='2020' and units=='1':key_number={30:31,31:33,32:34,33:35,34:36,35:38,36:39}.get(n,n)
    if key_number not in ks:
     report['held'].append({'source':paper,'number':n,'reason':'Matching guide question missing'});continue
    guide=ks[key_number]
    if q['marks']!=guide['marks']:
     report['held'].append({'source':paper,'number':n,'guideNumber':key_number,'reason':'Question/key marks disagree','paperMarks':q['marks'],'keyMarks':guide['marks']});continue
    answer,_=chemistry_content({**guide,'nodes':guide['prefix']+guide['nodes']},kpdf,kranges[key_number],True)
    if key_number!=n:
     report['corrections'].append({'source':paper,'number':n,'guideNumber':key_number,'reason':'Matched question by text; source marking guide uses different numbering'})
     answer='<p class="review-note">This question is numbered '+str(key_number)+' in the supplied marking guide.</p>'+answer
   digest=hashlib.sha256(signature(q).encode()).hexdigest()
   if digest in seen:
    report['duplicates'].append({'source':paper,'number':n,'sameAs':seen[digest]});continue
   identifier='chemistry-'+year+'-u'+units+'-q'+str(n);seen[digest]=identifier
   topics,subs=classify(q);plain=' '.join(text(x) for x in q['prefix']+q['nodes'])
   extra=[]
   if q['type']=='Multiple choice' and any(c['node'].tag in ['td','th','tr'] for c in q['choices']):
    q['nodes']=[node for node in q['nodes'] if not(node.tag=='table' and len(node.xpath('.//td[starts-with(normalize-space(.),\"(\")]|.//th[starts-with(normalize-space(.),\"(\")]'))>=4)]
   body,figs=chemistry_content({**q,'nativeOnly':identifier in ['chemistry-2020-u12-q9','chemistry-2021-u12-q21'],'nodes':q['prefix']+q['nodes']+extra},pdf,ranges[n])
   item={'id':identifier,'subject':'Chemistry','year':11,'examYear':int(year),'units':'1' if units=='1' else '1 & 2','number':n,'source':year+' Chemistry '+('Unit 1' if units=='1' else 'Units 1 & 2'),'section':q['type'],'types':[q['type']],'marks':q['marks'],'topics':topics,'subtopics':subs,'mixed':len(subs)>1,'text':plain,'figures':figs,'hash':digest,'html':body,'answerHtml':answer,'options':options,'optionsHtml':options_html,'correctAnswer':correct,'markingMode':'automatic' if options else 'manual','sourceFile':paper,'sourceKeyFile':key,'sourceKeyNumber':key_number,'bankFile':filename}
   batch.append(item);audits.append({'id':identifier,'type':q['type'],'topics':topics,'subtopics':subs,'text':plain[:400]})
  (OUT/filename).write_text(json.dumps(batch,ensure_ascii=False,separators=(',',':')));records+=batch
  report['exams'].append({'source':paper,'key':key,'questionsDetected':len(qs),'imported':len(batch),'mcqKeysVerified':len(answers),'file':filename});print(paper,len(batch),'/',len(qs),flush=True)
 index=[{k:v for k,v in q.items() if k not in ['html','answerHtml','options','optionsHtml','correctAnswer']} for q in records]
 (OUT/'practice-studio-chemistry-year11-index.json').write_text(json.dumps(index,ensure_ascii=False,separators=(',',':')))
 report['totalImported']=len(records);report['types']={t:sum(t in q['types'] for q in records) for t in ['Multiple choice','Short answer','Extended answer']}
 report['marking']='Multiple choice is checked against the supplied answer tables. Written questions use self-review with the original marking guide.'
 (OUT/'practice-studio-chemistry-year11-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));(WORK/'audit.json').write_text(json.dumps(audits,ensure_ascii=False,indent=2))
 print(json.dumps({k:v for k,v in report.items() if k!='exams'},indent=2))
if __name__=='__main__':main()
